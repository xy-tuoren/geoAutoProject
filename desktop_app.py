#!/usr/bin/env python3
"""Operator desktop panel for the Xiaohe ADB automation."""

from __future__ import annotations

import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
from pathlib import Path
from tkinter import BOTH, END, LEFT, RIGHT, VERTICAL, BooleanVar, StringVar, Tk, filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText
from typing import Any

from openpyxl import load_workbook

from paths import adb_binary, ensure_adb_env, is_frozen, writable_root

DEFAULT_OUTPUT = writable_root() / "captures"
ensure_adb_env()


def clean_questions(items: list[Any]) -> list[str]:
    """Trim, de-duplicate, and remove empty question rows."""
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        question = str(item).strip()
        if question and question not in seen:
            result.append(question)
            seen.add(question)
    return result


def load_question_file(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv"}:
        return clean_questions(path.read_text(encoding="utf-8-sig").splitlines())
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            payload = next((payload[key] for key in ("questions", "问题", "items", "data") if isinstance(payload.get(key), list)), [])
        if not isinstance(payload, list):
            raise ValueError("JSON 应为问题数组，或含 questions / 问题 数组的对象。")
        return clean_questions(payload)
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(value).strip().lower() if value is not None else "" for value in rows[0]]
        column = next((i for i, value in enumerate(header) if value in {"问题", "question", "questions", "提问"}), 0)
        start = 1 if any(header) else 0
        return clean_questions([row[column] if len(row) > column else "" for row in rows[start:]])
    raise ValueError("仅支持 TXT、CSV、JSON、XLSX、XLSM 文件。")


def adb_devices() -> list[str]:
    completed = subprocess.run(
        [str(adb_binary()), "devices", "-l"],
        text=True,
        capture_output=True,
        check=False,
    )
    return [line.split()[0] for line in completed.stdout.splitlines()[1:] if len(line.split()) >= 2 and line.split()[1] == "device"]


class OperatorApp:
    def __init__(self, root: Tk) -> None:
        self.root = root
        self.root.title("小荷 AI 医生 · ADB 自动化面板")
        self.root.minsize(880, 680)
        self.process: subprocess.Popen[str] | None = None
        self.events: queue.Queue[tuple[str, str]] = queue.Queue()
        self.temp_questions: Path | None = None

        self.device_var = StringVar()
        self.output_var = StringVar(value=str(DEFAULT_OUTPUT))
        self.timeout_var = StringVar(value="90")
        self.new_session_var = BooleanVar(value=False)
        self.dry_run_var = BooleanVar(value=False)
        self.status_var = StringVar(value="准备就绪。请连接手机并填写或导入问题。")
        self._build()
        self.refresh_devices()
        self.root.after(120, self._drain_events)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def _build(self) -> None:
        frame = ttk.Frame(self.root, padding=14)
        frame.pack(fill=BOTH, expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(3, weight=1)
        frame.rowconfigure(6, weight=1)

        ttk.Label(frame, text="设备").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        self.device_box = ttk.Combobox(frame, textvariable=self.device_var, state="readonly")
        self.device_box.grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Button(frame, text="刷新 ADB 设备", command=self.refresh_devices).grid(row=0, column=2, padx=(8, 0), pady=4)

        ttk.Label(frame, text="截图保存目录").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(frame, textvariable=self.output_var).grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(frame, text="选择目录", command=self.choose_output).grid(row=1, column=2, padx=(8, 0), pady=4)

        option_frame = ttk.Frame(frame)
        option_frame.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(8, 4))
        ttk.Label(option_frame, text="单题超时（秒）").pack(side=LEFT)
        ttk.Entry(option_frame, width=8, textvariable=self.timeout_var).pack(side=LEFT, padx=(6, 18))
        ttk.Checkbutton(option_frame, text="每题新建会话", variable=self.new_session_var).pack(side=LEFT, padx=(0, 18))
        ttk.Checkbutton(option_frame, text="仅输入测试（不发送）", variable=self.dry_run_var).pack(side=LEFT)
        ttk.Button(option_frame, text="导入问题文件", command=self.import_questions).pack(side=RIGHT)

        ttk.Label(frame, text="问题列表（每行一题，按显示顺序执行）").grid(row=3, column=0, columnspan=3, sticky="nw", pady=(8, 2))
        self.questions = ScrolledText(frame, height=14, wrap="word", font=("Arial", 12))
        self.questions.grid(row=4, column=0, columnspan=3, sticky="nsew")

        controls = ttk.Frame(frame)
        controls.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(10, 6))
        self.start_button = ttk.Button(controls, text="开始顺序执行并截图", command=self.start)
        self.start_button.pack(side=LEFT)
        self.stop_button = ttk.Button(controls, text="停止任务", command=self.stop, state="disabled")
        self.stop_button.pack(side=LEFT, padx=8)
        ttk.Label(controls, textvariable=self.status_var).pack(side=RIGHT)

        ttk.Label(frame, text="执行日志").grid(row=6, column=0, columnspan=3, sticky="sw", pady=(4, 2))
        self.log = ScrolledText(frame, height=11, state="disabled", wrap="word", font=("Menlo", 11))
        self.log.grid(row=7, column=0, columnspan=3, sticky="nsew")

    def append_log(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert(END, text)
        self.log.see(END)
        self.log.configure(state="disabled")

    def refresh_devices(self) -> None:
        try:
            devices = adb_devices()
        except FileNotFoundError as exc:
            messagebox.showerror("未找到 ADB", str(exc))
            return
        self.device_box["values"] = devices
        if devices:
            self.device_var.set(devices[0])
            self.status_var.set(f"已发现 {len(devices)} 台已授权设备。")
        else:
            self.device_var.set("")
            self.status_var.set("未发现已授权设备；请检查 USB 调试授权。")

    def choose_output(self) -> None:
        directory = filedialog.askdirectory(initialdir=self.output_var.get() or str(DEFAULT_OUTPUT))
        if directory:
            self.output_var.set(directory)

    def import_questions(self) -> None:
        filename = filedialog.askopenfilename(filetypes=[("问题文件", "*.txt *.csv *.json *.xlsx *.xlsm"), ("所有文件", "*.*")])
        if not filename:
            return
        try:
            imported = load_question_file(Path(filename))
        except Exception as exc:
            messagebox.showerror("导入失败", str(exc))
            return
        current = self.questions.get("1.0", END).splitlines()
        merged = clean_questions(current + imported)
        self.questions.delete("1.0", END)
        self.questions.insert("1.0", "\n".join(merged))
        self.status_var.set(f"已导入 {len(imported)} 条，当前共 {len(merged)} 条。")

    def start(self) -> None:
        questions = clean_questions(self.questions.get("1.0", END).splitlines())
        if not questions:
            messagebox.showwarning("没有问题", "请至少填写或导入一条问题。")
            return
        if not self.device_var.get():
            messagebox.showwarning("没有设备", "请先连接并刷新 ADB 设备。")
            return
        try:
            timeout = float(self.timeout_var.get())
            if timeout <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("超时设置无效", "请输入大于 0 的秒数。")
            return

        handle = tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".txt", prefix="xiaohe-questions-", delete=False)
        handle.write("\n".join(questions))
        handle.close()
        self.temp_questions = Path(handle.name)
        if is_frozen():
            command = [sys.executable, "--run-ask-xiaohe"]
        else:
            command = [sys.executable, str(Path(__file__).resolve().parent / "ask_xiaohe.py")]
        command += [
            "--file", str(self.temp_questions),
            "--serial", self.device_var.get(), "--output-dir", self.output_var.get(), "--timeout", str(timeout),
        ]
        if self.new_session_var.get():
            command.append("--new-session")
        if self.dry_run_var.get():
            command.append("--dry-run")

        self.append_log("\n$ " + " ".join(command) + "\n")
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self.status_var.set(f"正在顺序执行 {len(questions)} 条问题…")
        workdir = str(writable_root())
        self.process = subprocess.Popen(
            command,
            cwd=workdir,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            bufsize=1,
            env={**os.environ, "ADBUTILS_ADB_PATH": str(adb_binary())},
        )
        threading.Thread(target=self._read_process, daemon=True).start()

    def _read_process(self) -> None:
        assert self.process and self.process.stdout
        for line in self.process.stdout:
            self.events.put(("log", line))
        code = self.process.wait()
        self.events.put(("finished", str(code)))

    def _drain_events(self) -> None:
        try:
            while True:
                kind, value = self.events.get_nowait()
                if kind == "log":
                    self.append_log(value)
                elif kind == "finished":
                    self.process = None
                    self.start_button.configure(state="normal")
                    self.stop_button.configure(state="disabled")
                    self.status_var.set("执行完成。" if value == "0" else f"任务结束，退出码 {value}。请查看日志。")
                    if self.temp_questions:
                        self.temp_questions.unlink(missing_ok=True)
                        self.temp_questions = None
        except queue.Empty:
            pass
        self.root.after(120, self._drain_events)

    def stop(self) -> None:
        if self.process and self.process.poll() is None:
            self.process.terminate()
            self.status_var.set("正在停止任务…")

    def on_close(self) -> None:
        if self.process and self.process.poll() is None:
            if not messagebox.askyesno("任务正在运行", "关闭窗口将停止当前任务，是否继续？"):
                return
            self.process.terminate()
        self.root.destroy()


def main() -> None:
    root = Tk()
    OperatorApp(root)
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--run-ask-xiaohe":
        sys.argv = [sys.argv[0], *sys.argv[2:]]
        from ask_xiaohe import main as ask_main

        raise SystemExit(ask_main())
    main()
